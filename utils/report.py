"""Excel report generation for transactions."""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


@dataclass
class PeriodSummary:
    income: int
    expense: int
    profit: int
    tax_rate: int
    tax: int

    def render(self, period_label: str) -> str:
        return (
            f"📊 Отчёт за <b>{period_label}</b>:\n"
            f"• Доходы: {self.income:,} ₽\n"
            f"• Расходы: {self.expense:,} ₽\n"
            f"• Чистая прибыль: {self.profit:,} ₽\n"
            f"• Налог ({self.tax_rate}%): {self.tax:,} ₽"
        ).replace(",", " ")


def summarise(
    transactions: Iterable, *, tax_rate: int
) -> PeriodSummary:
    income = sum(int(t["amount"]) for t in transactions if t["type"] == "income")
    expense = sum(int(t["amount"]) for t in transactions if t["type"] == "expense")
    profit = income - expense
    tax = round(income * tax_rate / 100)
    return PeriodSummary(
        income=income, expense=expense, profit=profit, tax_rate=tax_rate, tax=tax
    )


def build_excel_report(
    transactions: Iterable,
    *,
    period_label: str,
    summary: PeriodSummary,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Транзакции"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="6A5ACD")
    headers = ["Дата", "Тип", "Сумма (₽)", "Категория", "Комментарий", "Запись"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for tx in transactions:
        created = tx["created_at"]
        if isinstance(created, datetime):
            created_str = created.strftime("%Y-%m-%d %H:%M")
        else:
            created_str = str(created)
        ws.append(
            [
                created_str,
                "Доход" if tx["type"] == "income" else "Расход",
                int(tx["amount"]),
                tx["category"] or "",
                tx["comment"] or "",
                tx["appointment_id"] or "",
            ]
        )

    # Width
    widths = [20, 10, 14, 18, 40, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(ord("A") + i - 1)].width = w

    # Summary block
    ws.append([])
    ws.append([f"Итог за {period_label}"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13)
    ws.append(["Доходы", summary.income])
    ws.append(["Расходы", summary.expense])
    ws.append(["Чистая прибыль", summary.profit])
    ws.append([f"Налог ({summary.tax_rate}%)", summary.tax])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
